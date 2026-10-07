"""Versioned local JAN master candidates. Candidate choice is always reviewed."""

import hashlib
import io
import re
import unicodedata
import zipfile
from datetime import date

from openpyxl import load_workbook

from .document_ai import DocumentError


def workbook(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if (
                len(archive.infolist()) > 500
                or sum(i.file_size for i in archive.infolist()) > 20_000_000
            ):
                raise ValueError()
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Exception:
        raise DocumentError("JANマスターは10 MB以内の通常のXLSXを選んでください") from None


def sheets(raw):
    book = workbook(raw)
    try:
        return book.sheetnames
    finally:
        book.close()


def parse_master(raw, sheet, effective_date):
    try:
        date.fromisoformat(effective_date)
    except (ValueError, TypeError):
        raise DocumentError("マスターの適用開始日を入力してください") from None
    book = workbook(raw)
    try:
        if sheet not in book.sheetnames:
            raise DocumentError("使用するシートを選んでください")
        tab = book[sheet]
        if tab.max_row > 10000 or tab.max_column > 50:
            raise DocumentError("マスターは10000行・50列以内にしてください")
        rows = list(tab.iter_rows())
        columns = []
        for line in rows[:10]:
            for i, cell in enumerate(line):
                if str(cell.value or "").replace(" ", "") in ("JAN", "JANコード") and i >= 2:
                    columns.append(i)
        columns = sorted(set(columns))
        if not columns:
            raise DocumentError("商品名・商品コード・JANコードの順の列が必要です")
        products = []
        invalid = 0
        for line in rows:
            for i in columns:
                cells = line[i - 2 : i + 1]
                if len(cells) != 3 or cells[2].value in (None, "", "JAN", "JANコード"):
                    continue
                name, code, jan = [str(c.value or "").strip() for c in cells]
                if any(c.data_type == "f" for c in cells) or not re.fullmatch(
                    r"[0-9]{8}|[0-9]{13}", jan
                ):
                    invalid += 1
                    continue
                digits = list(map(int, jan))
                if (sum(digits[-2::-2]) * 3 + sum(digits[-3::-2]) + digits[-1]) % 10 or not name:
                    invalid += 1
                    continue
                products.append(
                    {
                        "product_name": name,
                        "product_code": code,
                        "jan": jan,
                        "cell": cells[2].coordinate,
                    }
                )
        if not products:
            raise DocumentError("有効なJAN明細がありません")
        version = hashlib.sha256(raw + sheet.encode() + effective_date.encode()).hexdigest()
        return {
            "version": version,
            "source_sha": hashlib.sha256(raw).hexdigest(),
            "sheet": sheet,
            "effective_date": effective_date,
            "products": products,
            "invalid_rows": invalid,
        }
    finally:
        book.close()


def normalized(value):
    return re.sub(r"[\s・*＊\-ー]", "", unicodedata.normalize("NFKC", value)).casefold()


def candidates(master, row):
    code, name, jan = [
        str(row.get(k) or "").strip() for k in ("product_code", "product_name", "jan")
    ]
    result = []
    for product in master["products"]:
        match = (jan and product["jan"] == jan) or (code and product["product_code"] == code)
        label = normalized(product["product_name"])
        query = normalized(name)
        if match or (query and label and (query in label or label in query)):
            result.append(product)
    return result[:30]
